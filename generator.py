import os
import re
from datetime import datetime
from dateutil.relativedelta import relativedelta

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy
from docxcompose.composer import Composer

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# EOBI DEFAULTER NOTICE + ASSESSMENT GENERATOR
# ============================================================

DATE_FORMAT = "%Y-%m"
PERIOD_SEPARATOR = "، "


# ============================================================
# EOBI CONTRIBUTION RATES
# ============================================================

RATES = {
    2001: (3000, 170),
    2002: (3000, 170),
    2003: (3000, 170),
    2004: (3000, 170),
    2005: (3600, 210),
    2006: (4000, 280),
    2007: (4600, 322),
    2008: (6000, 360),
    2009: (6000, 360),
    2010: (7000, 420),
    2011: (7000, 420),
    2012: (8000, 480),
    2013: (10000, 600),
    2014: (12000, 720),
    2015: (13000, 780),
    2016: (14000, 840),
    2017: (15000, 900),
    2018: (16500, 990),
    2019: (17500, 1050),
    2020: (17500, 1050),
    2021: (20000, 1200),
    2022: (25000, 1500),
    2023: (32000, 1920),
    2024: (37000, 2220),
    2025: (40000, 2400),
    2026: (40000, 2400),
}


def get_rate_for_fy(year):
    if year in RATES:
        return RATES[year]

    if year < 2001:
        return (3000, 170)

    return (40000, 2400)


# ============================================================
# GENERAL HELPERS
# ============================================================

def format_date(val):
    if pd.isna(val):
        return None

    if isinstance(val, (datetime, pd.Timestamp)):
        return val.strftime(DATE_FORMAT)

    s = str(val).strip()

    return s if s and s.lower() != "nat" else None


def convert_to_datetime(value):
    """
    Convert Excel/CSV date values into Python datetime.
    """
    if pd.isna(value):
        return None

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if isinstance(value, datetime):
        return value

    try:
        return pd.to_datetime(value).to_pydatetime()
    except Exception:
        return None


# ============================================================
# PERIOD FUNCTIONS
# ============================================================

def build_period_string(row):
    periods = []
    seen = set()

    for i in range(1, 5):

        frm = format_date(row.get(f"From{i}"))
        to_ = format_date(row.get(f"To{i}"))

        if frm and to_:
            p = f"{frm} تا {to_}"

        elif frm:
            p = frm

        elif to_:
            p = to_

        else:
            continue

        if p not in seen:
            seen.add(p)
            periods.append(p)

    return PERIOD_SEPARATOR.join(periods) if periods else "________"


def get_unique_periods(row):

    periods = []
    seen = set()

    for i in range(1, 5):

        frm = convert_to_datetime(row.get(f"From{i}"))
        to_ = convert_to_datetime(row.get(f"To{i}"))

        if frm is None or to_ is None:
            continue

        key = (
            frm.strftime("%Y-%m-%d"),
            to_.strftime("%Y-%m-%d")
        )

        if key in seen:
            continue

        seen.add(key)
        periods.append((frm, to_))

    return periods


def month_start(dt):
    return dt.replace(day=1)


def add_months(dt, months):
    return dt + relativedelta(months=months)


def months_between(start, end):
    return (
        (end.year - start.year) * 12
        + (end.month - start.month)
        + 1
    )


def split_into_fy(from_date, to_date):

    start = month_start(from_date)

    if to_date.day == 1:
        end = add_months(month_start(to_date), -1)
    else:
        end = month_start(to_date)

    if end < start:
        return []

    periods = []

    current = start

    while current <= end:

        # July to June financial year
        if current.month >= 7:

            fy_year = current.year
            fy_end = datetime(
                current.year + 1,
                6,
                1
            )

        else:

            fy_year = current.year - 1
            fy_end = datetime(
                current.year,
                6,
                1
            )

        period_end = min(fy_end, end)

        month_count = months_between(
            current,
            period_end
        )

        if month_count > 0:

            wages, rate = get_rate_for_fy(fy_year)

            periods.append({
                "from_str": current.strftime("%b-%y"),
                "to_str": period_end.strftime("%b-%y"),
                "months": month_count,
                "wages": wages,
                "rate": rate,
            })

        current = add_months(
            period_end,
            1
        )

    return periods


# ============================================================
# WORD NOTICE FUNCTIONS
# ============================================================

def replace_in_paragraph(
    paragraph,
    placeholder,
    value
):

    if placeholder not in paragraph.text:
        return False

    # Normal case: placeholder exists entirely inside a run
    for run in paragraph.runs:

        if placeholder in (run.text or ""):

            run.text = run.text.replace(
                placeholder,
                value
            )

            return True

    # Fallback if Word split the placeholder
    # across multiple runs
    full_text = paragraph.text

    new_text = full_text.replace(
        placeholder,
        value
    )

    if paragraph.runs:

        paragraph.runs[0].text = new_text

        for run in paragraph.runs[1:]:
            run.text = ""

    else:
        paragraph.add_run(new_text)

    return True


def get_all_paragraphs(doc):

    paragraphs = list(doc.paragraphs)

    for table in doc.tables:

        for row in table.rows:

            for cell in row.cells:

                paragraphs.extend(
                    cell.paragraphs
                )

    return paragraphs


def fill_notice(
    template_path,
    mapping
):

    doc = Document(template_path)

    paragraphs = get_all_paragraphs(doc)

    for paragraph in paragraphs:

        for key, value in mapping.items():

            placeholder = f"«{key}»"

            if placeholder in paragraph.text:

                replace_in_paragraph(
                    paragraph,
                    placeholder,
                    "" if value is None else str(value)
                )

    return doc


def remove_bookmarks(element):

    for tag in (
        "bookmarkStart",
        "bookmarkEnd"
    ):

        for node in element.findall(
            ".//" + qn(f"w:{tag}")
        ):

            parent = node.getparent()

            if parent is not None:
                parent.remove(node)


# ============================================================
# ASSESSMENT SHEET
# ============================================================

thin = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

header_font = Font(
    name="Times New Roman",
    bold=True,
    size=14
)

title_font = Font(
    name="Times New Roman",
    bold=True,
    size=12
)

normal_font = Font(
    name="Times New Roman",
    size=10
)

bold_font = Font(
    name="Times New Roman",
    bold=True,
    size=10
)

center = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

left_align = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True
)


def safe_value(row, column, default=""):

    value = row.get(column, default)

    if pd.isna(value):
        return default

    return value


def get_ips(row):

    value = safe_value(
        row,
        "IPS",
        0
    )

    try:
        return int(float(value))
    except Exception:
        return 0


def get_reg_no(row):

    code = safe_value(
        row,
        "Code",
        ""
    )

    sub_code = safe_value(
        row,
        "Sub_Code",
        ""
    )

    code = str(code).strip()

    try:

        if sub_code != "":
            sub_code = str(
                int(float(sub_code))
            )
        else:
            sub_code = ""

    except Exception:
        sub_code = str(
            sub_code
        ).strip()

    if sub_code:
        return f"{code} {sub_code}"

    return code


def create_assessment_sheet(
    workbook,
    row_data,
    sheet_name
):

    ws = workbook.create_sheet(
        title=sheet_name[:31]
    )

    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    ws.merge_cells("A1:K1")

    ws["A1"] = (
        "EMPLOYEES' OLD-AGE BENEFITS INSTITUTION"
    )

    ws["A1"].font = header_font
    ws["A1"].alignment = center

    ws.merge_cells("A2:K2")

    ws["A2"] = (
        "REGIONAL OFFICE, FAISALABAD CENTRAL"
    )

    ws["A2"].font = title_font
    ws["A2"].alignment = center

    ws.merge_cells("A3:K3")

    ws["A3"] = ""

    ws.merge_cells("A4:K4")

    ws["A4"] = (
        "DEFAULTERS ASSESSMENT SHEET"
    )

    ws["A4"].font = Font(
        name="Times New Roman",
        bold=True,
        size=13
    )

    ws["A4"].alignment = center

    # --------------------------------------------------------
    # EMPLOYER INFORMATION
    # --------------------------------------------------------

    ws.merge_cells("A6:B6")

    ws["A6"] = "Name of Employer:"
    ws["A6"].font = bold_font

    ws.merge_cells("C6:H6")

    ws["C6"] = str(
        safe_value(
            row_data,
            "Name of the Establishment",
            ""
        )
    )

    ws["C6"].font = Font(
        name="Times New Roman",
        bold=True,
        size=12
    )

    ws["C6"].alignment = center

    ws["I6"] = "Reg. No."
    ws["I6"].font = bold_font
    ws["I6"].alignment = center

    ws.merge_cells("J6:K6")

    ws["J6"] = get_reg_no(row_data)

    ws["J6"].font = Font(
        name="Times New Roman",
        bold=True,
        size=11
    )

    ws["J6"].alignment = center

    # --------------------------------------------------------
    # DEFAULT PERIODS
    # --------------------------------------------------------

    unique_periods = get_unique_periods(
        row_data
    )

    period_texts = [
        f"{frm.strftime('%b-%Y')} to "
        f"{to_.strftime('%b-%Y')}"
        for frm, to_ in unique_periods
    ]

    ws.merge_cells("A7:K7")

    ws["A7"] = (
        "Default Period(s):  "
        +
        (
            "  |  ".join(period_texts)
            if period_texts
            else "—"
        )
    )

    ws["A7"].font = bold_font
    ws["A7"].alignment = left_align

    # --------------------------------------------------------
    # COMPUTATION TITLE
    # --------------------------------------------------------

    ws.merge_cells("A8:K8")

    ws["A8"] = (
        "Detailed Computation of Payable Contribution "
        "(Annexure to VR-001)"
    )

    ws["A8"].font = Font(
        name="Times New Roman",
        bold=True,
        size=10,
        italic=True
    )

    ws["A8"].alignment = center

    # --------------------------------------------------------
    # TABLE HEADERS
    # --------------------------------------------------------

    headers = [
        (1, 1, "From"),
        (2, 2, "To"),
        (3, 3, "No. of\nMonths"),
        (4, 4, "No. IPs"),
        (5, 5, "Minimum\nWages"),
        (6, 6, "Contribution\nRate"),
        (
            7,
            7,
            "Assessed Amount\nof Contribution"
        ),
        (
            8,
            8,
            "Paid Amount of\nContribution"
        ),
        (
            9,
            9,
            "Principal Amount\nPayable"
        ),
        (
            10,
            11,
            "Remarks"
        ),
    ]

    for start, end, text in headers:

        if start != end:

            ws.merge_cells(
                start_row=9,
                start_column=start,
                end_row=9,
                end_column=end
            )

        cell = ws.cell(
            row=9,
            column=start,
            value=text
        )

        cell.font = bold_font
        cell.alignment = center

        for col in range(
            start,
            end + 1
        ):

            ws.cell(
                row=9,
                column=col
            ).border = thin

            ws.cell(
                row=9,
                column=col
            ).font = bold_font

            ws.cell(
                row=9,
                column=col
            ).alignment = center

    # --------------------------------------------------------
    # CALCULATIONS
    # --------------------------------------------------------

    ips = get_ips(row_data)

    all_splits = []

    for frm, to_ in unique_periods:

        all_splits.extend(
            split_into_fy(
                frm,
                to_
            )
        )

    all_splits.sort(
        key=lambda x:
        datetime.strptime(
            x["from_str"],
            "%b-%y"
        )
    )

    data_start = 10

    total_months = 0
    total_assessed = 0

    for i, period in enumerate(
        all_splits
    ):

        row_num = data_start + i

        ws.cell(
            row=row_num,
            column=1,
            value=period["from_str"]
        ).alignment = center

        ws.cell(
            row=row_num,
            column=2,
            value=period["to_str"]
        ).alignment = center

        ws.cell(
            row=row_num,
            column=3,
            value=period["months"]
        ).alignment = center

        ws.cell(
            row=row_num,
            column=4,
            value=ips
        ).alignment = center

        ws.cell(
            row=row_num,
            column=5,
            value=period["wages"]
        ).alignment = center

        ws.cell(
            row=row_num,
            column=5
        ).number_format = "#,##0"

        ws.cell(
            row=row_num,
            column=6,
            value=period["rate"]
        ).alignment = center

        ws.cell(
            row=row_num,
            column=6
        ).number_format = "#,##0"

        assessed = (
            period["rate"]
            * ips
            * period["months"]
        )

        ws.cell(
            row=row_num,
            column=7,
            value=assessed
        ).alignment = center

        ws.cell(
            row=row_num,
            column=7
        ).number_format = "#,##0"

        # Paid contribution
        ws.cell(
            row=row_num,
            column=8,
            value=""
        ).alignment = center

        # Principal payable
        ws.cell(
            row=row_num,
            column=9,
            value=assessed
        ).alignment = center

        ws.cell(
            row=row_num,
            column=9
        ).number_format = "#,##0"

        ws.merge_cells(
            start_row=row_num,
            start_column=10,
            end_row=row_num,
            end_column=11
        )

        for col in range(1, 12):

            ws.cell(
                row=row_num,
                column=col
            ).border = thin

            ws.cell(
                row=row_num,
                column=col
            ).font = normal_font

        total_months += period["months"]
        total_assessed += assessed

    # --------------------------------------------------------
    # SUBTOTAL
    # --------------------------------------------------------

    last_data = (
        data_start
        + len(all_splits)
        - 1
        if all_splits
        else data_start - 1
    )

    subtotal_row = last_data + 1

    ws.merge_cells(
        start_row=subtotal_row,
        start_column=1,
        end_row=subtotal_row,
        end_column=2
    )

    ws.cell(
        row=subtotal_row,
        column=1,
        value="Sub Total"
    ).font = bold_font

    ws.cell(
        row=subtotal_row,
        column=1
    ).alignment = center

    ws.cell(
        row=subtotal_row,
        column=3,
        value=total_months
    ).font = bold_font

    ws.cell(
        row=subtotal_row,
        column=3
    ).alignment = center

    ws.cell(
        row=subtotal_row,
        column=7,
        value=total_assessed
    ).font = bold_font

    ws.cell(
        row=subtotal_row,
        column=7
    ).number_format = "#,##0"

    ws.cell(
        row=subtotal_row,
        column=7
    ).alignment = center

    ws.cell(
        row=subtotal_row,
        column=9,
        value=total_assessed
    ).font = bold_font

    ws.cell(
        row=subtotal_row,
        column=9
    ).number_format = "#,##0"

    ws.cell(
        row=subtotal_row,
        column=9
    ).alignment = center

    for col in range(1, 12):

        ws.cell(
            row=subtotal_row,
            column=col
        ).border = thin

    # --------------------------------------------------------
    # STATUTORY INCREASE
    # --------------------------------------------------------

    statutory_row = subtotal_row + 1

    ws.merge_cells(
        start_row=statutory_row,
        start_column=1,
        end_row=statutory_row,
        end_column=8
    )

    ws.cell(
        row=statutory_row,
        column=1,
        value="Statutory Increase @ 50% Approx"
    ).font = bold_font

    statutory_amount = int(
        total_assessed * 0.5
    )

    ws.cell(
        row=statutory_row,
        column=9,
        value=statutory_amount
    ).font = bold_font

    ws.cell(
        row=statutory_row,
        column=9
    ).number_format = "#,##0"

    ws.cell(
        row=statutory_row,
        column=9
    ).alignment = center

    for col in range(1, 12):

        ws.cell(
            row=statutory_row,
            column=col
        ).border = thin

    # --------------------------------------------------------
    # TOTAL
    # --------------------------------------------------------

    total_row = statutory_row + 1

    ws.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=8
    )

    ws.cell(
        row=total_row,
        column=1,
        value="Total"
    ).font = Font(
        name="Times New Roman",
        bold=True,
        size=11
    )

    ws.cell(
        row=total_row,
        column=1
    ).alignment = center

    final_total = (
        total_assessed
        + statutory_amount
    )

    ws.cell(
        row=total_row,
        column=9,
        value=final_total
    ).font = Font(
        name="Times New Roman",
        bold=True,
        size=11
    )

    ws.cell(
        row=total_row,
        column=9
    ).number_format = "#,##0"

    ws.cell(
        row=total_row,
        column=9
    ).alignment = center

    for col in range(1, 12):

        ws.cell(
            row=total_row,
            column=col
        ).border = thin

    # --------------------------------------------------------
    # NOTES
    # --------------------------------------------------------

    note_row = total_row + 2

    ws.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=11
    )

    ws.cell(
        row=note_row,
        column=1,
        value=(
            "Above calculations are based on no. of "
            "Insured Persons as mentioned in "
            "PR-01/Last paid PR-03 Slip."
        )
    ).font = Font(
        name="Times New Roman",
        size=9,
        italic=True
    )

    ps_row = note_row + 1

    ws.cell(
        row=ps_row,
        column=1,
        value="P.S."
    ).font = bold_font

    notes = [

        "1) Please note that in case of non-payment of above contribution, record verification through VR-006 shall be initiated resulting in issuance of Demand Notice VR-008 which will entail recovery under Land Revenue Act, 1967.",

        "2) Please also note that upon issuance of VR-008, you will be liable to pay Statutory Increase @ 2% per month upto maximum of 50% on total payable contribution. However SI is taken as 50% for this calculation and its final percentage will be fixed at the time of issuance of VR007 & VR008.",

        "3) This is simply a minimum contribution liability and is not a final / assessed demand issued without prejudice to any liability that may be ascertained on the basis of record verification (VR-006)."
    ]

    for i, note in enumerate(notes):

        row_num = ps_row + 1 + i

        ws.merge_cells(
            start_row=row_num,
            start_column=1,
            end_row=row_num,
            end_column=11
        )

        ws.cell(
            row=row_num,
            column=1,
            value=note
        ).font = Font(
            name="Times New Roman",
            size=8
        )

        ws.cell(
            row=row_num,
            column=1
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

        ws.row_dimensions[
            row_num
        ].height = 30

    # --------------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------------

    widths = {
        1: 11,
        2: 11,
        3: 10,
        4: 9,
        5: 12,
        6: 13,
        7: 16,
        8: 14,
        9: 15,
        10: 10,
        11: 10
    }

    for col, width in widths.items():

        ws.column_dimensions[
            get_column_letter(col)
        ].width = width

    # --------------------------------------------------------
    # PAGE SETUP
    # --------------------------------------------------------

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    ws.print_options.horizontalCentered = True


# ============================================================
# DATA VALIDATION
# ============================================================

REQUIRED_NOTICE_COLUMNS = [
    "Name",
    "Name of the Establishment",
    "Address",
    "City",
    "Contact",
    "Code",
    "Sub_Code",
]

REQUIRED_ASSESSMENT_COLUMNS = [
    "Name of the Establishment",
    "Code",
    "Sub_Code",
    "IPS",
]


def validate_columns(df):

    missing_notice = [
        col
        for col in REQUIRED_NOTICE_COLUMNS
        if col not in df.columns
    ]

    missing_assessment = [
        col
        for col in REQUIRED_ASSESSMENT_COLUMNS
        if col not in df.columns
    ]

    missing = list(
        dict.fromkeys(
            missing_notice
            + missing_assessment
        )
    )

    return missing


# ============================================================
# MAIN GENERATOR FUNCTION
# ============================================================

def generate_all(
    data_file,
    notice_template,
    output_dir
):
    """
    Main function called by Streamlit.

    Parameters
    ----------
    data_file:
        Path to Excel or CSV data file.

    notice_template:
        Path to Urdu Word notice template.

    output_dir:
        Directory where generated files will be placed.

    Returns
    -------
    dict
        Paths of generated files and summary information.
    """

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    # --------------------------------------------------------
    # READ DATA
    # --------------------------------------------------------

    extension = os.path.splitext(
        data_file
    )[1].lower()

    if extension == ".csv":

        df = pd.read_csv(
            data_file
        )

    elif extension in [".xlsx", ".xls"]:

        df = pd.read_excel(
            data_file
        )

    else:

        raise ValueError(
            "Unsupported data file. "
            "Please upload XLSX, XLS or CSV."
        )

    # --------------------------------------------------------
    # VALIDATE COLUMNS
    # --------------------------------------------------------

    missing_columns = validate_columns(df)

    if missing_columns:

        raise ValueError(
            "The following required columns "
            "are missing from your data file:\n\n"
            + "\n".join(
                f"• {column}"
                for column in missing_columns
            )
        )

    if len(df) == 0:

        raise ValueError(
            "The uploaded data file contains "
            "no records."
        )

    # --------------------------------------------------------
    # OUTPUT DIRECTORIES
    # --------------------------------------------------------

    individual_dir = os.path.join(
        output_dir,
        "generated_notices"
    )

    os.makedirs(
        individual_dir,
        exist_ok=True
    )

    # --------------------------------------------------------
    # WORKBOOK
    # --------------------------------------------------------

    workbook = Workbook()

    if "Sheet" in workbook.sheetnames:

        del workbook["Sheet"]

    individual_files = []

    # --------------------------------------------------------
    # PROCESS EACH ESTABLISHMENT
    # --------------------------------------------------------

    for index, row in df.iterrows():

        establishment_name = str(
            safe_value(
                row,
                "Name of the Establishment",
                f"Establishment_{index + 1}"
            )
        ).strip()

        # ----------------------------------------------------
        # NOTICE DATA
        # ----------------------------------------------------

        data = {}

        field_map = {

            "Name":
                "Name",

            "Name of the Establishment":
                "Name_of_the_Establishment",

            "Address":
                "Address",

            "City":
                "City",

            "Contact":
                "contact",

            "Code":
                "Code",

            "Sub_Code":
                "Sub_Code",
        }

        for excel_column, placeholder in field_map.items():

            value = row.get(
                excel_column,
                ""
            )

            if pd.isna(value):

                value = ""

            data[placeholder] = value

        data["Period"] = build_period_string(
            row
        )

        # ----------------------------------------------------
        # CREATE NOTICE
        # ----------------------------------------------------

        document = fill_notice(
            notice_template,
            data
        )

        # ----------------------------------------------------
        # SAFE FILE NAME
        # ----------------------------------------------------

        safe_name = re.sub(
            r"[^\w\s-]",
            "",
            establishment_name
        ).strip()

        safe_name = safe_name[:50]

        if not safe_name:

            safe_name = (
                f"row_{index + 1}"
            )

        notice_filename = (
            f"{index + 1:02d}_"
            f"{safe_name}.docx"
        )

        notice_path = os.path.join(
            individual_dir,
            notice_filename
        )

        document.save(
            notice_path
        )

        individual_files.append(
            notice_path
        )

        # ----------------------------------------------------
        # ASSESSMENT SHEET
        # ----------------------------------------------------

        clean_sheet_name = re.sub(
            r'[\\/*?:\[\]]',
            "",
            establishment_name
        )

        sheet_name = (
            f"{index + 1:02d}_"
            f"{clean_sheet_name[:25]}"
        )

        create_assessment_sheet(
            workbook,
            row,
            sheet_name
        )

    # --------------------------------------------------------
    # COMBINED NOTICE DOCUMENT
    # --------------------------------------------------------

    combined_notice_path = os.path.join(
        output_dir,
        "All_Payment_Notices.docx"
    )

    if individual_files:

        master = Document(
            individual_files[0]
        )

        composer = Composer(
            master
        )

        for file_path in individual_files[1:]:

            composer.append(
                Document(file_path)
            )

        composer.save(
            combined_notice_path
        )

    # --------------------------------------------------------
    # ASSESSMENT WORKBOOK
    # --------------------------------------------------------

    assessment_path = os.path.join(
        output_dir,
        "Defaulters_Assessment_Sheets.xlsx"
    )

    workbook.save(
        assessment_path
    )

    # --------------------------------------------------------
    # RETURN RESULTS
    # --------------------------------------------------------

    return {

        "individual_dir":
            individual_dir,

        "individual_files":
            individual_files,

        "combined_notice":
            combined_notice_path,

        "assessment":
            assessment_path,

        "records":
            len(df),

        "dataframe":
            df,
    }